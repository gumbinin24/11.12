import openpyxl
from docx import Document


def main():
    wb = openpyxl.load_workbook("Ученики.xlsx")
    sheet = wb.active
    doc = Document()

    for row in sheet.iter_rows(min_row=2):
        school = row[2].value
        name = row[1].value
        surname = row[0].value
        doc.add_paragraph(f"Награждается ученик {school} школы по имени {surname} {name}")
        doc.add_page_break()



    doc.save("Награды.docx")


if __name__ == "__main__":
    main()